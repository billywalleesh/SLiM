"""Three populations, three toolkits - every element a native PowerPoint shape.

The variation a population can hold scales with Ne (theta = 4*Ne*mu), so
Ne 1,000 / 500 / 100 carry 10 / 5 / 1 tools. Each tool is a different shape,
because the point is not the quantity of variation but how many DIFFERENT
answers a population still has.

Slide 1  the three toolkits
Slide 2  a problem arrives that needs one particular tool
"""
import os, math
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

A     = "http://schemas.openxmlformats.org/drawingml/2006/main"
S1    = RGBColor(0x2A, 0x78, 0xD6)
S2    = RGBColor(0xEB, 0x68, 0x34)
S3    = RGBColor(0x1B, 0xAF, 0x7A)
INK   = RGBColor(0x0B, 0x0B, 0x0B)
INK2  = RGBColor(0x52, 0x51, 0x4E)
FAINT = RGBColor(0xF2, 0xF1, 0xEE)
GRID  = RGBColor(0xDE, 0xDC, 0xD7)
GHOST = RGBColor(0xDE, 0xDC, 0xD7)

TOOLS = [MSO_SHAPE.GEAR_6, MSO_SHAPE.STAR_5_POINT, MSO_SHAPE.HEXAGON,
         MSO_SHAPE.DIAMOND, MSO_SHAPE.REGULAR_PENTAGON,
         MSO_SHAPE.ISOSCELES_TRIANGLE, MSO_SHAPE.MATH_PLUS,
         MSO_SHAPE.LIGHTNING_BOLT, MSO_SHAPE.OCTAGON, MSO_SHAPE.CHEVRON]
NEEDED = 2                      # the hexagon: Ne 1,000 and 500 have it, 100 does not

POPS = [("1,000", 10), ("500", 5), ("100", 1)]
CXS  = [6.6, 17.0, 27.4]
R_POP, R_RING, TOOL = 2.2, 4.0, 1.22


def plain(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)
    spPr = sh._element.spPr
    for e in spPr.findall(qn('a:effectLst')):
        spPr.remove(e)
    etree.SubElement(spPr, f"{{{A}}}effectLst")
    return sh


def shape(shapes, kind, cx, cy, w, h, fill, line=None, lw=1.0):
    sh = shapes.add_shape(kind, Cm(cx - w/2), Cm(cy - h/2), Cm(w), Cm(h))
    plain(sh)
    if fill is None: sh.fill.background()
    else:            sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else:            sh.line.color.rgb = line; sh.line.width = Pt(lw)
    return sh


def text(shapes, cx, cy, w, h, runs, align=PP_ALIGN.CENTER):
    """runs: list of (text, size, colour, bold, baseline) - baseline 'sub' shrinks."""
    tb = shapes.add_textbox(Cm(cx - w/2), Cm(cy - h/2), Cm(w), Cm(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = None
    for item in runs:
        if item is None:                      # paragraph break
            para = tf.add_paragraph(); para.alignment = align; continue
        txt, size, col, bold = item[:4]
        sub = len(item) > 4 and item[4] == "sub"
        if para is None:
            para = tf.paragraphs[0]; para.alignment = align
        r = para.add_run(); r.text = txt
        r.font.name = "Arial"; r.font.size = Pt(size)
        r.font.color.rgb = col; r.font.bold = bold
        if sub:
            r.font._rPr.set("baseline", "-25000")
    return tb


def toolkit(shapes, cx, CY, label, n, dim_others=False, mark=None,
            show_count=True):
    grp = shapes.add_group_shape(); g = grp.shapes
    shape(g, MSO_SHAPE.OVAL, cx, CY, R_POP*2, R_POP*2, FAINT, GRID, 1.5)
    text(g, cx, CY, R_POP*1.9, 2.0,
         [("N", 21, INK, True), ("e", 21, INK, True, "sub"), None,
          (label, 24, INK, True)])

    for k in range(n):
        ang = math.radians(-90 + k * 360.0 / max(n, 1))
        tx, ty = cx + R_RING*math.cos(ang), CY + R_RING*math.sin(ang)
        hot = (mark is not None and k == NEEDED)
        col = S2 if hot else (GHOST if dim_others else S1)
        sh = shape(g, TOOLS[k], tx, ty, TOOL, TOOL, col)
        if hot:
            shape(g, MSO_SHAPE.OVAL, tx, ty, TOOL*1.85, TOOL*1.85, None, S2, 2.0)

    if show_count:
        text(g, cx, CY + R_RING + 1.5, 9.0, 1.0,
             [(f"{n} tool{'s' if n > 1 else ''}", 17,
               S1 if n > 1 else S2, True)])
    return grp


# built on the UZH template so the titles use the deck's own placeholder
import copy
SRC = "/home/claude/rev2/deck.pptx"
prs = Presentation(SRC)
_lst = prs.slides._sldIdLst
for _sld in list(_lst):
    _rId = _sld.get("{http://schemas.openxmlformats.org/officeDocument/2006/"
                    "relationships}id")
    prs.part.drop_rel(_rId); _lst.remove(_sld)
prs.save("_tmp.pptx"); prs = Presentation("_tmp.pptx")
LAYOUT = next(l for l in prs.slide_masters[0].slide_layouts
              if l.name == "Nur Titel")

# ---------------------------------------------------------------- slide 1
s = prs.slides.add_slide(LAYOUT)
s.shapes.title.text = "Effective size is a toolkit"
CY1 = 9.6
for (label, n), cx in zip(POPS, CXS):
    toolkit(s.shapes, cx, CY1, label, n)
text(s.shapes, 16.93, 16.5, 31.0, 1.0,
     [("The variation a population can hold scales with N", 15, INK2, False),
      ("e", 15, INK2, False, "sub"),
      (" — ten times the effective size, roughly ten times the variants "
       "it still carries.", 15, INK2, False)])
text(s.shapes, 16.93, 17.9, 31.0, 1.0,
     [("A problem you have not met yet needs a tool you already have.",
       18, INK, True)])

# ---------------------------------------------------------------- slide 2
s2 = prs.slides.add_slide(LAYOUT)
s2.shapes.title.text = "Then a problem arrives"

banner = s2.shapes.add_group_shape(); bg = banner.shapes
shape(bg, MSO_SHAPE.ROUNDED_RECTANGLE, 16.93, 4.3, 15.0, 1.8, None, S2, 1.8)
text(bg, 14.6, 4.3, 9.4, 1.2,
     [("it can only be met with this tool", 15, INK2, False)], PP_ALIGN.RIGHT)
shape(bg, TOOLS[NEEDED], 21.4, 4.3, 1.25, 1.25, S2)

CY2 = 10.2
for (label, n), cx in zip(POPS, CXS):
    toolkit(s2.shapes, cx, CY2, label, n, dim_others=True, mark=True,
            show_count=False)
    has = n > NEEDED
    text(s2.shapes, cx, CY2 + R_RING + 1.6, 9.0, 1.0,
         [("has it" if has else "does not have it", 17,
           S3 if has else S2, True)])

text(s2.shapes, 16.93, 17.9, 31.0, 1.0,
     [("Selection can only use variation that is already there — and the "
       "threshold was set without asking what the problem would be.",
       17, INK, True)])

prs.save("population_toolkits.pptx")
os.remove("_tmp.pptx")
print("saved")
