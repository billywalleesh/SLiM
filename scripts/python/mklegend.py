"""The map's colourbar, rebuilt as native PowerPoint shapes.

Matches 14_map_mean_annual_temp.R: ramp #2166AC - #F7F7F7 - #B2182B, data range
0.28-14.19 degC, white at 7.24, warm end at the top, breaks at 5 and 10, white
tick marks inside the bar.
"""
import os
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
INK  = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)

LO, HI, MIDT = 0.28, 14.19, 7.24        # the map's data range and white point
BREAKS = [5, 10]

BAR_W, BAR_H = 0.70, 9.50               # cm
X0, Y0 = 1.20, 3.60                     # top-left of the bar on the slide


def plain(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)
    spPr = sh._element.spPr
    for e in spPr.findall(qn('a:effectLst')):
        spPr.remove(e)
    etree.SubElement(spPr, f"{{{A}}}effectLst")
    return sh


def gradient_bar(shapes, x, y, w, h):
    """Three-stop linear gradient, warm at the top."""
    sh = shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    plain(sh)
    sh.line.fill.background()
    spPr = sh._element.spPr
    for tag in ("a:solidFill", "a:gradFill", "a:noFill"):
        for e in spPr.findall(qn(tag)):
            spPr.remove(e)

    grad = etree.Element(f"{{{A}}}gradFill")
    grad.set("flip", "none"); grad.set("rotWithShape", "1")
    gsLst = etree.SubElement(grad, f"{{{A}}}gsLst")
    # position 0 is where the gradient starts; ang 270 deg makes that the bottom
    for pos, hexv in ((0, "2166AC"), (50000, "F7F7F7"), (100000, "B2182B")):
        gs = etree.SubElement(gsLst, f"{{{A}}}gs"); gs.set("pos", str(pos))
        clr = etree.SubElement(gs, f"{{{A}}}srgbClr"); clr.set("val", hexv)
    lin = etree.SubElement(grad, f"{{{A}}}lin")
    lin.set("ang", str(270 * 60000)); lin.set("scaled", "0")

    ln = spPr.find(qn('a:ln'))
    spPr.insert(list(spPr).index(ln) if ln is not None else len(spPr), grad)
    return sh


def textbox(shapes, x, y, w, h, text, size, colour, align=PP_ALIGN.LEFT,
            bold=False):
    tb = shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = "Arial"; r.font.size = Pt(size); r.font.color.rgb = colour
    r.font.bold = bold
    return tb


def build(shapes, numbered=True):
    frac = lambda t: (t - LO) / (HI - LO)          # 0 at the bottom of the bar

    gradient_bar(shapes, X0, Y0, BAR_W, BAR_H)

    # thin frame, so the pale middle of the ramp still has an edge
    fr = shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(X0), Cm(Y0), Cm(BAR_W), Cm(BAR_H))
    plain(fr); fr.fill.background()
    fr.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF); fr.line.width = Pt(0.75)

    textbox(shapes, X0 - 0.1, Y0 - 1.15, 5.0, 0.8, "Temperature (°C)",
            13, INK)

    if numbered:
        for t in BREAKS:
            ty = Y0 + BAR_H * (1 - frac(t))
            tick = shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(X0), Cm(ty - 0.035),
                                    Cm(BAR_W), Cm(0.07))
            plain(tick); tick.line.fill.background()
            tick.fill.solid(); tick.fill.fore_color.rgb = RGBColor(255, 255, 255)
            textbox(shapes, X0 + BAR_W + 0.25, ty - 0.35, 2.0, 0.7, str(t),
                    12.5, INK2)
    else:
        textbox(shapes, X0 + BAR_W + 0.25, Y0 - 0.05, 3.0, 0.7, "warmer",
                12.5, INK2)
        textbox(shapes, X0 + BAR_W + 0.25, Y0 + BAR_H - 0.65, 3.0, 0.7,
                "colder", 12.5, INK2)


def gradient_bar_h(shapes, x, y, w, h):
    """Same three stops, running left (cold) to right (warm)."""
    sh = shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    plain(sh); sh.line.fill.background()
    spPr = sh._element.spPr
    for tag in ("a:solidFill", "a:gradFill", "a:noFill"):
        for e in spPr.findall(qn(tag)):
            spPr.remove(e)
    grad = etree.Element(f"{{{A}}}gradFill")
    grad.set("flip", "none"); grad.set("rotWithShape", "1")
    gsLst = etree.SubElement(grad, f"{{{A}}}gsLst")
    for pos, hexv in ((0, "2166AC"), (50000, "F7F7F7"), (100000, "B2182B")):
        gs = etree.SubElement(gsLst, f"{{{A}}}gs"); gs.set("pos", str(pos))
        clr = etree.SubElement(gs, f"{{{A}}}srgbClr"); clr.set("val", hexv)
    lin = etree.SubElement(grad, f"{{{A}}}lin")
    lin.set("ang", "0"); lin.set("scaled", "0")
    ln = spPr.find(qn('a:ln'))
    spPr.insert(list(spPr).index(ln) if ln is not None else len(spPr), grad)
    return sh


def build_h(shapes, labels=("P1", "P2", "P3", "Px"),
            x=1.2, y=4.0, w=18.0, h=0.80):
    """A key, not an axis: the scale laid out flat with each box bracketed."""
    gradient_bar_h(shapes, x, y, w, h)
    fr = shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    plain(fr); fr.fill.background()
    fr.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF); fr.line.width = Pt(0.75)

    for i in range(1, 4):                      # the three interior boundaries
        tx = x + w * i / 4
        tick = shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(tx - 0.035),
                                Cm(y - 0.18), Cm(0.07), Cm(h + 0.36))
        plain(tick); tick.line.fill.background()
        tick.fill.solid(); tick.fill.fore_color.rgb = INK

    for i, lab in enumerate(labels):
        textbox(shapes, x + w * i / 4, y + h + 0.30, w / 4, 0.70, lab,
                12.5, INK2, align=PP_ALIGN.CENTER, bold=True)

    textbox(shapes, x, y - 0.95, 4.0, 0.65, "colder", 12, INK2)
    textbox(shapes, x + w - 4.0, y - 0.95, 4.0, 0.65, "warmer", 12, INK2,
            align=PP_ALIGN.RIGHT)


prs = Presentation()
prs.slide_width = Cm(33.87); prs.slide_height = Cm(19.05)
for numbered, note in ((True,  "Exact replica of the map legend — use it only "
                               "if the boxes carry real temperatures."),
                       (False, "Same bar, no numbers — use it with the "
                               "schematic P1–P4 boxes.")):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    grp = s.shapes.add_group_shape()          # one draggable object
    build(grp.shapes, numbered)
    textbox(s.shapes, 6.5, 8.0, 24.0, 1.0, note, 13, INK2)

s = prs.slides.add_slide(prs.slide_layouts[6])
grp = s.shapes.add_group_shape()
build_h(grp.shapes)
textbox(s.shapes, 1.2, 7.5, 28.0, 1.0,
        "Horizontal key \u2014 put this under the 2 \u00d7 2. A bar beside a grid "
        "implies up/down means temperature; a bar underneath, with each box "
        "bracketed, only says where each box sits.", 13, INK2)

prs.save("map_legend.pptx")
print("saved")
